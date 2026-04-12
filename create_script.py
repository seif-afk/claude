from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Set default font
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(11)

# Title
title = doc.add_heading('', level=1)
run = title.add_run('ZoomInfo vs Outbound Agency:\nWhich Is Better For Your Branded Merch Company?')
run.font.size = Pt(22)
run.font.color.rgb = RGBColor(0, 0, 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('YouTube Video Script  |  Netswick')
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(128, 128, 128)

doc.add_paragraph('')

# Helper functions
def add_section_header(text):
    h = doc.add_heading('', level=2)
    run = h.add_run(text)
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0, 0, 0)

def add_timestamp(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(150, 150, 150)
    run.font.italic = True

def add_body(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(8)
    return p

# ============ HOOK ============
add_section_header('HOOK')
add_timestamp('[0:00 - 0:30]')

add_body("We've booked almost 500 sales meetings for branded merch companies in the past three months using cold email outreach. We also helped a promo company close a Fortune 100 deal. And we've gotten leads from companies like Stripe, Uber, and Hilton for our clients.")

add_body("I've also personally used ZoomInfo. I've had clients come to me after spending $15,000 to $30,000 a year on it and getting almost nothing.")

add_body("So today I'm going to break down ZoomInfo versus hiring an outbound agency, specifically for branded merch companies. Which one actually makes sense for you? Because the answer depends a lot on your size and where you're at.")

add_body("Let's get into it.")

# ============ WHAT ZOOMINFO GIVES YOU ============
add_section_header('What ZoomInfo Actually Gives You')
add_timestamp('[0:30 - 3:00]')

add_body("So ZoomInfo is a B2B contact database. That's it. You pay for access to millions of contacts, filter by industry, company size, job title, location, and you export a spreadsheet of people to reach out to.")

add_body("And that's where it ends.")

add_body("You get names, emails, phone numbers. Some of them are accurate. A lot of them aren't. People change jobs constantly. Emails get deactivated. Your reps end up calling numbers that don't work and emailing people who left the company six months ago.")

add_body("I had a prospect tell me he spent $10,000 on a service that used ZoomInfo data. He went through the list and one of the contacts was literally dead. Another had sold their business years ago. Half the list was useless before they even started.")

add_body("Another one, a promo company in Georgia, sent 5,000 emails using ZoomInfo contacts and booked three meetings. Three. From five thousand emails. And that's not even counting what they paid for ZoomInfo on top of the time their team spent setting everything up.")

add_body("So the data itself, even when it's decent, is just a starting point. It doesn't send the emails for you. It doesn't write the messaging. It doesn't find out which companies need promo right now versus eventually. It doesn't call leads when they respond. It doesn't book the meeting.")

add_body("You're paying premium pricing for a spreadsheet. The rest is on you.")

# ============ WHAT AN OUTBOUND AGENCY DOES ============
add_section_header('What An Outbound Agency Actually Does')
add_timestamp('[3:00 - 7:00]')

add_body("An outbound agency, at least a good one, handles the entire thing from research to booked meeting.")

add_body("So what does that look like?")

add_body("First, the research. We're not just pulling contacts from a database. We're analyzing your past client data, figuring out which industries spent the most with you, what the buying patterns look like, what use cases drove those purchases. Was it trade shows? Employee onboarding kits? Corporate gifting? We find that out.")

add_body("Then we find the trending products in those industries. What are tech companies buying right now? What are construction companies buying? Because when you reach out offering a specific Carhartt jacket for a construction company's new hire program, that gets a response. When you reach out saying \"we do branded merch,\" that gets deleted.")

add_body("Then the prospect list. And this is where it gets interesting. We use tools that are cheaper and better than ZoomInfo. Prospeo has 230 million contacts. Apollo is another one. Ocean.io can find companies that look exactly like your best clients based on their actual website content. Data quality is comparable or better and costs a fraction of what you'd pay ZoomInfo.")

add_body("But we go further. We run AI agents across the entire list to find which companies have upcoming events, which ones are hiring, which ones just got funded, which ones are rebranding. So now we're not blasting 10,000 people. We're reaching the 500 that actually have a reason to buy promo in the next 30 to 60 days.")

add_body("Then the infrastructure. We set up 50 plus domains, a hundred inboxes, warm everything up for two weeks so that when we launch, emails hit the primary inbox, not spam. We use a dedicated cold email tool built for this. Not HubSpot. Not Outlook.")

add_body("Then messaging. I can't give you the exact copy we use, but the principle is the same across all our clients. Lead with a specific product idea. Tie it to their use case. Reference their buying signal. Position your company as a creative partner who brings ideas, not a vendor begging for a catalog review.")

add_body("Then when someone responds, we call them. Within minutes. We book the meeting. We send them case studies and a catalog before the call so they're warmed up and ready to have a real conversation by the time your salesperson gets on.")

add_body("Your salesperson just shows up to a meeting with someone who already knows your company, has seen your work, and has a real need. That's a completely different conversation than cold calling a ZoomInfo contact who doesn't pick up.")

# ============ THE REAL GAP ============
add_section_header('Why The Data Alone Doesn\'t Work')
add_timestamp('[7:00 - 9:00]')

add_body("I talk to branded merch companies every single week. Probably five to ten a week at this point. And the ones who tried ZoomInfo or similar tools almost always describe the same experience.")

add_body("They export a list. Plug it into HubSpot or Outlook or sometimes literally a Word mail merge. Send the same email to everyone. Something like \"we do branded merch, want to see our catalog?\" And nothing happens.")

add_body("And they blame outbound. They say cold email doesn't work in this industry. It's all relationships. People already have vendors.")

add_body("But that's not what happened. What happened is they took a raw list with no buying signals, sent a generic message with no specific offer, from an email infrastructure that wasn't built for cold outreach, and it all landed in spam. Or if it didn't land in spam, nobody cared because the message didn't speak to anything they needed.")

add_body("Having 10,000 contacts means nothing if you don't have a proper email infrastructure that can send at volume without burning your domain. You need messaging that positions you as a creative partner. You need buying signals that tell you which of those 10,000 actually need promo right now. And you need someone replying to interested leads within five minutes and booking the meeting while they're still warm.")

add_body("ZoomInfo gives you step one out of ten. The agency gives you all ten.")

# ============ YOUR TEAM CAN'T DO THIS ============
add_section_header('Why Your Sales Team Can\'t Fill The Gap')
add_timestamp('[9:00 - 11:00]')

add_body("I know what some of you are thinking. Okay, so I'll buy the cheaper data tools myself, learn Smartlead, set up the domains, write the copy, and just have my sales team run it.")

add_body("Your sales reps have become account managers. I say this on every call and every single prospect agrees with me. They service existing clients, handle reorders, follow up on quotes. At some point they stopped reaching out to anyone new.")

add_body("It makes sense from their side. Managing accounts is comfortable. Commissions still come in. Nobody's tracking how many new conversations they're starting each week.")

add_body("Prospecting in 2026 is a completely different skill set. AI agents that find buying signals across thousands of companies. Email infrastructure with dozens of domains and warm-up protocols. Parallel dialers. Clay workflows. This stuff changes every few months. Your reps aren't going to learn it. They shouldn't have to. That's not their job.")

add_body("Your marketing team is the other option people think of. But they're busy with the website, social media, trade shows, catalogs. They don't have the bandwidth to build and manage a cold outreach system that requires daily monitoring and constant optimization.")

add_body("You can try to hire for it. Good luck finding someone who understands AI tools, cold email deliverability, AND the branded merch industry. I've spent years building a team that can do this and it's still not easy.")

# ============ WHICH ONE FOR YOUR SIZE ============
add_section_header('Which One Is Right For Your Company')
add_timestamp('[11:00 - 14:00]')

add_body("Alright so this is where I want to be genuinely fair because it's not the same answer for everyone.")

add_body("If you're under two million in revenue, and you have one or two people who are willing to grind, you can probably get by with a cheaper database and doing outreach yourself. Apollo, Prospeo, either one works. Learn a tool like Smartlead. Set up your own infrastructure. Write your own copy. You'll make mistakes, it'll take a few months to figure out, but if you can't invest in a partner yet, that's the path. Just cancel ZoomInfo. You're overpaying for what you're getting.")

add_body("If you're between two million and roughly a hundred million, this is where an outbound agency wins and it's not even close. And yes I'm biased because that's exactly what we do. But the logic is pretty straightforward.")

add_body("You're big enough that the investment makes sense. You have real revenue, real clients, a real reputation. But you're not so big that you have fifteen people on your go-to-market team who can build this internally. Your reps are managing accounts. Your marketing team is stretched. And every month you don't have a system bringing in new business, you're just depending on referrals and hoping the phone rings.")

add_body("A specialized outbound agency already knows what works in your space. We don't need three months to test. We know which industries respond. We know what products to offer. We know what buying signals matter. We can plug you into a proven system and get you results from the first week. We did it for a client recently, 19 qualified leads in their first week. Another one, 15 hospital meetings in three weeks. This stuff just works when the system is dialed in.")

add_body("If you're above a hundred million, you probably already have internal marketing teams, demand gen people, maybe SDR teams. You might still benefit from a partner for specific campaigns or to test new markets. We work with companies like Swag.com and Imprint Engine which are not small. They chose to work with us because of the specialization. We only do branded merch outbound. That expertise is hard to replicate internally even with a big team. But at that level you have options that a ten million dollar company doesn't.")

# ============ HOW TO PICK AN AGENCY ============
add_section_header('If You Go The Agency Route, Here\'s What To Look For')
add_timestamp('[14:00 - 16:00]')

add_body("If you decide a partner is the right move, please don't just pick the first one that cold calls you. Most lead gen agencies are terrible. I say that as someone who runs one.")

add_body("They should specialize in your industry. If they work with dentists, lawyers, SaaS companies, and also promo, they don't understand your market. They're going to spend months testing angles that someone who lives in this space already knows don't work. An agency that only does promo gets you results from week one because they already ran the same campaign for other companies like yours.")

add_body("They should show you real case studies. Not \"we sent 50,000 emails.\" Nobody cares about that. How many qualified meetings did you book? What companies showed up on the call? What did they end up closing? Those are the numbers that matter.")

add_body("They should handle the full funnel. If they're just giving you a list and saying good luck, that's ZoomInfo with a markup. You want someone who does the research, builds the infrastructure, writes the messaging, handles the replies, calls the leads, books the meetings, and nurtures before the call. The whole thing.")

add_body("And they should offer a short enough commitment that you can see results before you've spent a fortune. If someone wants twelve months upfront before you've seen a single meeting, walk away.")

# ============ OUTRO ============
add_section_header('OUTRO')
add_timestamp('[16:00 - 16:30]')

add_body("So that's the full breakdown. ZoomInfo gives you contacts. An outbound agency gives you meetings. For branded merch companies in that two to hundred million range, it's really not a debate.")

add_body("If you want my team to build and run this for your company, there's a link in the description to book a call and we'll walk you through how it would work for your specific market.")

add_body("Subscribe for more breakdowns like this. Peace.")

# Save
doc.save('/home/user/claude/scripts/zoominfo-vs-outbound-agency-script.docx')
print("Done")
